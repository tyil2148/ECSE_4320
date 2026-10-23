#!/usr/bin/env python3
"""Build an editable Word report from report/report.md (no data substitution)."""
import argparse, pathlib, re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=pathlib.Path(__file__).resolve().parents[1]

def math_run(text):
    r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);return r

def math_sub(base,sub):
    node=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(math_run(base));s=OxmlElement('m:sub');s.append(math_run(sub));node.extend([e,s]);return node

def equation(doc,text):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para=OxmlElement('m:oMathPara');math=OxmlElement('m:oMath')
    if text.startswith('I ='):
        math.append(math_run('I = '));frac=OxmlElement('m:f')
        n=OxmlElement('m:num');n.append(math_run('F'));d=OxmlElement('m:den');d.append(math_run('Q'));frac.extend([n,d]);math.append(frac)
    elif text.startswith('Pbound'):
        math.append(math_sub('P','bound'));math.append(math_run(' = min('));math.append(math_sub('P','peak'));math.append(math_run(', B × I)'))
    elif text.startswith('Ppeak'):
        math.append(math_sub('P','peak'));math.append(math_run(' = f × U × L × 2'))
    else:math.append(math_run(text))
    para.append(math);p._p.append(para)

def prose(doc,text,style=None):
    p=doc.add_paragraph(style=style)
    # Filenames/options stay literal, with code marked in monospace if backticked.
    for i,part in enumerate(re.split(r'(`[^`]+`)',text)):
        r=p.add_run(part[1:-1] if part.startswith('`') and part.endswith('`') else part)
        if part.startswith('`'):r.font.name='Consolas';r.font.size=Pt(10)
    if text.startswith('[ADD '):
        for r in p.runs:r.font.color.rgb=RGBColor.from_string('595959');r.italic=True
    return p

def table(doc,rows):
    rows=[[v.strip() for v in row.strip('|').split('|')] for row in rows]
    rows=[r for r in rows if not all(re.fullmatch(r'[:\- ]+',c) for c in r)]
    columns=len(rows[0]);tbl=doc.add_table(rows=1,cols=columns);tbl.alignment=WD_TABLE_ALIGNMENT.CENTER;tbl.autofit=False
    widths={2:[1.7,5.2],4:[1.05,2.25,1.3,2.3],5:[1.35,1.6,.75,1.5,1.7]}[columns]
    for c,w in zip(tbl.columns,widths):c.width=Inches(w)
    for idx,data in enumerate(rows):
        cells=tbl.rows[0].cells if idx==0 else tbl.add_row().cells
        for col,(cell,value) in enumerate(zip(cells,data)):
            cell.width=Inches(widths[col]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for side in ('top','left','bottom','right','insideH','insideV'):
                edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
            pr.append(borders);margins=OxmlElement('w:tcMar')
            for side in ('top','bottom','left','right'):
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1);p.paragraph_format.line_spacing=1
            r=p.add_run(value);r.font.size=Pt(10)
            if idx==0:
                r.bold=True;r.font.color.rgb=RGBColor(0,0,0);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6F1');pr.append(shade)
            elif columns==5 and col>=2 and len(value)<28:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        trpr=tbl.rows[idx]._tr.get_or_add_trPr();split=OxmlElement('w:cantSplit');trpr.append(split)
        if idx==0:repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--source',type=pathlib.Path,default=ROOT/'report/report.md');p.add_argument('--output',type=pathlib.Path,default=ROOT/'report/report.docx');a=p.parse_args()
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
    normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(11);normal.font.color.rgb=RGBColor(0,0,0)
    normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.04
    for name,size in [('Title',25),('Heading 1',15),('Heading 2',12)]:
        s=doc.styles[name];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0);s.paragraph_format.space_before=Pt(10);s.paragraph_format.space_after=Pt(6)
    # Remove any inherited paragraph borders, including template title rules.
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    # Page numbers aid review of a multi-page experiment report.
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
    run=OxmlElement('w:r');prop=OxmlElement('w:rPr');size=OxmlElement('w:sz');size.set(qn('w:val'),'18');prop.append(size);run.append(prop)
    text=OxmlElement('w:t');text.text='1';run.append(text);field.append(run);footer._p.append(field)
    doc.core_properties.title='SIMD Advantage Profiling';doc.core_properties.subject='Advanced Computer Systems Project 1';doc.core_properties.author='';doc.core_properties.comments=''
    lines=a.source.read_text().splitlines();i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line=='<!-- PAGE -->':doc.add_page_break();i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):rows.append(lines[i].strip());i+=1
            table(doc,rows);continue
        if line.startswith('# '):doc.add_paragraph(line[2:],'Title')
        elif line.startswith('## '):doc.add_paragraph(line[3:],'Heading 1')
        elif line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2')
        elif line.startswith('EQ: '):equation(doc,line[4:])
        elif re.fullmatch(r'!\[.*\]\(.*\)',line):
            m=re.fullmatch(r'!\[(.*)\]\((.*)\)',line)
            image=(a.source.parent/m.group(2)).resolve()
            if not image.is_file():raise FileNotFoundError(image)
            doc.add_picture(str(image),width=Inches(6.9));prose(doc,m.group(1),'Caption')
        else:
            para=prose(doc,line)
            if line.startswith('[1]') or re.match(r'^\[[2-6]\]',line):
                for r in para.runs:r.font.size=Pt(9)
                para.paragraph_format.space_after=Pt(4)
        i+=1
    a.output.parent.mkdir(parents=True,exist_ok=True);doc.save(a.output);print(a.output)
if __name__=='__main__':main()
