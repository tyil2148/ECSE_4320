#!/usr/bin/env python3
"""Optional: rebuild the editable report from the included Markdown source.
Install requirements-report.txt first. Editing the DOCX directly is also supported.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path(__file__).resolve().parents[1]

def build():
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.7)
    sec.left_margin=sec.right_margin=Inches(.75)
    styles=doc.styles
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
        s=styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
    for border in list(styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    styles['Normal'].font.size=Pt(11)
    styles['Normal'].paragraph_format.space_after=Pt(6)
    styles['Normal'].paragraph_format.line_spacing=1.05
    styles['Title'].font.size=Pt(25);styles['Title'].paragraph_format.space_after=Pt(7)
    for name,size in [('Heading 1',15),('Heading 2',12),('Heading 3',11)]:
        styles[name].font.size=Pt(size);styles[name].font.bold=True
        styles[name].paragraph_format.space_before=Pt(10);styles[name].paragraph_format.space_after=Pt(5)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=footer.add_run('Advanced Computer Systems  |  ');r.font.size=Pt(9)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    lines=(ROOT/'report/Cache_Memory_Report.md').read_text().splitlines();i=0;pending_break=False
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line=='[[PAGEBREAK]]':pending_break=True;i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                vals=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',v) for v in vals):rows.append(vals)
                i+=1
            n=len(rows[0]);table=doc.add_table(rows=0,cols=n);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
            widths={2:[2.0,5.0],5:[1.4,1.4,1.6,1.8,.8]}.get(n,[7/n]*n)
            # The intensity table has no compact n column.
            if 'Traffic mix' in rows[0][0]:widths=[1.2,1.4,1.2,1.5,1.7]
            for col,width in zip(table.columns,widths):col.width=Inches(width)
            for ridx,row in enumerate(rows):
                cells=table.add_row().cells
                for j,(cell,value) in enumerate(zip(cells,row)):
                    cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cell.text=value
                    p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);p.paragraph_format.line_spacing=1.0
                    for run in p.runs:run.font.size=Pt(10);run.bold=ridx==0;run.font.color.rgb=RGBColor(0,0,0)
                    tcPr=cell._tc.get_or_add_tcPr()
                    borders=OxmlElement('w:tcBorders')
                    for edge in ['top','left','bottom','right']:
                        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                    tcPr.append(borders)
                    margins=OxmlElement('w:tcMar')
                    for side in ['top','left','bottom','right']:
                        e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
                    tcPr.append(margins)
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6F1' if ridx==0 else 'FFFFFF');tcPr.append(shade)
                trPr=table.rows[-1]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
                if ridx==0:trPr.append(OxmlElement('w:tblHeader'))
            doc.add_paragraph().paragraph_format.space_after=Pt(0)
            continue
        if line.startswith('# '):p=doc.add_paragraph(line[2:],style='Title')
        elif line.startswith('## '):p=doc.add_paragraph(line[3:],style='Heading 1')
        elif line.startswith('### '):p=doc.add_paragraph(line[4:],style='Heading 2')
        else:
            p=doc.add_paragraph(line)
            if line.startswith('Figure '):
                for r in p.runs:r.italic=True
            if line.startswith('[INSERT'):
                for r in p.runs:r.font.color.rgb=RGBColor(80,80,80)
        if pending_break:
            p.paragraph_format.page_break_before=True;pending_break=False
        i+=1
    doc.core_properties.title='Cache and Memory Performance Profiling'
    doc.core_properties.subject='Advanced Computer Systems Project 2'
    doc.core_properties.author=''
    output=ROOT/'report/Cache_Memory_Report.docx';doc.save(output);print(output)
if __name__=='__main__':build()
