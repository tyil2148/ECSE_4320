#!/usr/bin/env python3
"""Create an editable report, optionally inserting measured tables and figures.

Interpretation remains the student's work. Existing documents are never overwritten.
"""
from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


class Report:
    def __init__(self, results=None):
        self.doc=Document()
        # The bundled default template may carry decorative title borders.
        for element in self.doc.styles.element.xpath(".//w:pBdr"):
            element.getparent().remove(element)
        self.md=[]
        self.results=Path(results) if results else None
        self.rows=self.csv("summary.csv")
        section=self.doc.sections[0]
        section.page_width=Inches(8.5); section.page_height=Inches(11)
        section.top_margin=section.bottom_margin=Inches(.7)
        section.left_margin=section.right_margin=Inches(.75)
        section.header_distance=section.footer_distance=Inches(.28)
        for name,size in (("Normal",10.5),("Title",24),("Subtitle",11),("Heading 1",14),("Heading 2",11.5),("Caption",9)):
            st=self.doc.styles[name]
            st.font.name="Arial"; st.font.size=Pt(size); st.font.color.rgb=RGBColor(0,0,0)
            st.paragraph_format.space_after=Pt(5)
            st.paragraph_format.line_spacing=1.08
            if name.startswith("Heading"):
                st.paragraph_format.space_before=Pt(10)
                st.paragraph_format.keep_with_next=True
            if name=="Title": st.paragraph_format.space_after=Pt(8)
        header=section.header.paragraphs[0]
        header.text="SSD Performance Profiling"
        header.runs[0].font.size=Pt(8); header.runs[0].font.color.rgb=RGBColor(0,0,0)
        footer=section.footer.paragraphs[0]
        footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run("Page ").font.size=Pt(8)
        field=OxmlElement("w:fldSimple"); field.set(qn("w:instr"),"PAGE"); footer._p.append(field)
        self.doc.core_properties.title="SSD Performance Profiling"
        self.doc.core_properties.subject="Advanced Computer Systems Project 3"
        self.doc.core_properties.author=""
        self.doc.core_properties.keywords="SSD fio WSL performance queue depth tail latency"

    def csv(self,name):
        p=self.results/"analysis"/name if self.results else None
        if not p or not p.exists(): return []
        with p.open(newline="") as f: return list(csv.DictReader(f))

    def p(self,text,italic=False):
        paragraph=self.doc.add_paragraph()
        run=paragraph.add_run(text); run.italic=italic
        if italic: run.font.size=Pt(9.5)
        self.md.extend([text,""])
        return paragraph

    def h(self,text,level=1):
        self.doc.add_heading(text,level)
        self.md.extend(["#"*(level+1)+" "+text,""])

    def page(self): self.doc.add_page_break()

    def table(self,headers,rows,widths=None,font=9):
        table=self.doc.add_table(rows=1,cols=len(headers))
        table.alignment=WD_TABLE_ALIGNMENT.CENTER
        table.autofit=False
        widths=widths or [7/len(headers)]*len(headers)
        for col,width in zip(table.columns,widths): col.width=Inches(width)
        for cell,text,width in zip(table.rows[0].cells,headers,widths):
            cell.text=str(text); cell.width=Inches(width)
            shade=OxmlElement("w:shd"); shade.set(qn("w:fill"),"DCE6F1"); cell._tc.get_or_add_tcPr().append(shade)
            for run in cell.paragraphs[0].runs: run.bold=True
        repeat=OxmlElement("w:tblHeader"); table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            cells=table.add_row().cells
            for cell,text,width in zip(cells,row,widths):
                cell.text=str(text); cell.width=Inches(width)
        for row_index,row in enumerate(table.rows):
            no_split=OxmlElement("w:cantSplit"); row._tr.get_or_add_trPr().append(no_split)
            for j,cell in enumerate(row.cells):
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr=cell._tc.get_or_add_tcPr()
                borders=OxmlElement("w:tcBorders")
                for edge in ("top","left","bottom","right"):
                    e=OxmlElement("w:"+edge); e.set(qn("w:val"),"single"); e.set(qn("w:sz"),"4"); e.set(qn("w:color"),"D9D9D9"); borders.append(e)
                pr.append(borders)
                margins=OxmlElement("w:tcMar")
                for edge in ("top","left","bottom","right"):
                    e=OxmlElement("w:"+edge); e.set(qn("w:w"),"85"); e.set(qn("w:type"),"dxa"); margins.append(e)
                pr.append(margins)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(1); p.paragraph_format.space_before=Pt(1)
                    p.paragraph_format.line_spacing=1.05
                    text=cell.text
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or (len(text)>28 and row_index>0) else WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs: run.font.size=Pt(font); run.font.color.rgb=RGBColor(0,0,0)
        self.md.append("| "+" | ".join(headers)+" |")
        self.md.append("| "+" | ".join("---" for _ in headers)+" |")
        self.md.extend("| "+" | ".join(str(v).replace("\n"," ") for v in row)+" |" for row in rows)
        self.md.append("")
        self.doc.add_paragraph().paragraph_format.space_after=Pt(0)

    def fig(self,stem,caption):
        path=self.results/"plots"/(stem+".png") if self.results else None
        if path and path.exists():
            p=self.doc.add_paragraph(); p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(path),width=Inches(6.7))
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            self.doc.add_paragraph(caption,style="Caption")
            self.md.extend([f"![{caption}]({path.absolute()})",""])
        else:
            self.p("[Insert "+stem+".png] "+caption,italic=True)

    def find(self,family,pattern,read,role=None):
        return next((r for r in self.rows if r["family"]==family and r["direction"]=="total" and r["pattern"]==pattern and int(r["read_pct"])==read and (role is None or r["role"]==role)),None)

    @staticmethod
    def number(row,key,sd=False):
        if not row: return "TBD"
        if sd and int(row["n"])>1: return f"{float(row[key+'_mean']):.2f} +/- {float(row[key+'_sd']):.2f}"
        return f"{float(row[key+'_mean']):.2f}"


def build(output,results=None):
    r=Report(results)
    d=r.doc
    d.add_paragraph("SSD Performance Profiling",style="Title")
    d.add_paragraph("Advanced Computer Systems Project 3",style="Subtitle")
    r.md.extend(["# SSD Performance Profiling","","Advanced Computer Systems Project 3",""])
    r.p("Author  [TBD]     Course section  [TBD]     Measurement date  [TBD]")
    r.h("1 Objective and scope")
    r.p("This project characterizes how block size, access pattern, read/write mix, queue depth, worker count, and address range affect storage throughput and latency. It measures QD 1 baselines, maximum observed small-block IOPS and sequential bandwidth, the throughput versus latency tradeoff, and completion-latency tails. The experiment also tests a warm-page-cache pitfall and observes performance during sustained writes.")
    r.p("The target is a Windows 11 laptop using WSL2. The measurements characterize its complete virtualized storage path. The physical SSD model and interface must be established from host metadata before drawing device-specific conclusions.")
    r.h("System and storage metadata",2)
    r.table(["Item","Configuration"],[
        ["Processor","AMD Ryzen 5 PRO 7530U with Radeon Graphics; advertised base 2.00 GHz"],
        ["Host operating system","Windows 11 x64; build [TBD]"],
        ["WSL environment","WSL version, Linux distribution, kernel, visible CPUs, memory and swap [TBD]"],
        ["Physical SSD","Model, firmware, capacity, bus and link details [TBD]"],
        ["Test path and filesystem","[TBD]; planned owned file under ~/ssdbench-data on Linux ext4"],
        ["Storage mapping","Windows volume hosting the distribution VHDX and its physical SSD [TBD]"],
        ["Capacity and state","Test range [TBD]; host and guest free capacity before/after [TBD]"],
        ["Power and temperature","AC power, Windows power mode, cooling, SSD temperature or unavailable [TBD]"]
    ],[1.55,5.45],font=9)
    r.h("Safety and initialization",2)
    r.p("The runner creates only its own disposable regular file in a new or empty directory. It refuses device paths, symbolic links, hard links, unmanaged files, and changed file identities. One full sequential direct-I/O write with a final flush initializes the tested range before reads. No raw partition, trim, or secure-erase operation is used. This is file preparation, not full-drive preconditioning.")

    r.page(); r.h("2 Experimental design")
    r.p("The full profile uses the matrix below. The planned total file size is 16 GiB. Each ordinary run has a 3-second ramp and 15 measured seconds; each case is repeated three times. Tail runs have 45 measured seconds, and the state observation lasts 120 seconds with no ramp. The actual values and any overrides are retained in manifest.json.")
    r.table(["Experiment","Varied setting","Fixed settings","Cases"],[
        ["QD 1 baseline","Read/write; random/sequential","4 KiB random or 128 KiB sequential; one worker",4],
        ["Block size","4, 16, 32, 64, 128, 256 KiB; both patterns; read/write","QD 16; one worker",24],
        ["Read/write mix","100/0, 70/30, 50/50, 0/100","4 KiB random or 128 KiB sequential; QD 16",8],
        ["Queue depth","1, 2, 4, 8, 16, 32, 64, 128, 256","Four baseline workloads; one worker",36],
        ["Workers","1, 2, 4, 8 workers","Aggregate QD 16 or 64; random 4 KiB read/write",16],
        ["Address range","1/16, 1/4, full file","Random 4 KiB read/write; QD 16",6],
        ["Cache pitfall","Direct or buffered","Warmed range; random 4 KiB reads; psync; QD 1",2],
        ["State observation","Elapsed time in write window","Sequential 128 KiB write; QD 32",1],
        ["Adaptive tails","Middle and near-knee QD","Four baseline workloads; one worker",8]
    ],[1.15,2.25,3.1,.5],font=8.8)
    r.p("The matrix has 97 static cases and eight adaptive cases, giving 315 measured runs at three trials. The default schedule takes about 122 minutes plus initialization, process startup, and cache warming. It is write intensive; the plan displays write exposure before execution. A smoke profile checks operation with fewer settings and does not establish full rubric coverage.")
    r.h("Controls and execution order",2)
    r.p("Main experiments use libaio, direct=1, aligned fixed-size transfers, and fio-generated incompressible buffers. Workers use disjoint regions while holding the total address range constant. The pitfall uses psync for both controls to isolate the direct-I/O setting. Payload refill can consume CPU and is part of this workload. Writes measure completion without per-operation fsync.")
    r.p("Each trial visits the experiment families in the table order, with deterministically shuffled cases within each static family. Tail depths are selected after all queue trials. Trial-specific random seeds and actual execution order are recorded. A two-second pause separates measured runs. Shared drive history, prior writes, and incomplete thermal recovery remain possible sources of dependence.")

    r.page(); r.h("3 Code structure and reproducibility")
    r.table(["Component","Responsibility"],[
        ["scripts/ssdbench.py","Entry point for inspect, init, plan, run, analyze, plot, and cleanup"],
        ["ssdbench/plan.py","Case matrix, alignment, disjoint worker ranges, and fio configuration generation"],
        ["ssdbench/safety.py","Ownership marker, file identity checks, locking, and guarded deletion"],
        ["ssdbench/runner.py","Initialization, trials, exact commands, failures, checkpoints, resume, and adaptive tails"],
        ["ssdbench/analysis.py","Unit conversion, directional metrics, merged histograms, trial statistics, knee rule and coverage"],
        ["ssdbench/plotting.py","Labeled figures with units and repeated-trial SD bars where applicable"],
        ["ssdbench/environment.py","Guest kernel, CPU, mount, memory, disk-counter and available temperature records"],
        ["scripts/collect_host.ps1","Read-only Windows disk, power, free-space and optional reliability metadata"],
        ["scripts/build_report.py","Editable report generation and optional insertion of measured tables and figures"],
        ["tests/test_project.py","Safety, matrix completeness, aggregation, quantiles, statistics and knee selection checks"]
    ],[2.05,4.95],font=9)
    r.h("Workflow",2)
    r.p("Inspect the WSL mount and host storage mapping, fill operator notes, and review the generated plan. Initialize the disposable file once. Check the smoke profile, then run the full profile once; that command already includes repeated trials. Preserve the output directory, regenerate analysis and figures as needed, and add the empirical interpretation to this report. Resume uses the same command plus --resume and verifies settings, source hashes, fio version, target identity, and retained raw hashes. The Windows collector uses the disk and reliability interfaces documented by Microsoft [4, 5].")
    r.h("Data hygiene",2)
    r.p("Each attempt retains its exact fio configuration, unmodified JSON+ output, stderr, auxiliary stdout, command, seed, elapsed time, and SHA256 hash. The journal records actual order and status. Before/after snapshots retain available guest temperatures, memory and disk counters. Failed attempts stay archived and are excluded from analysis. The latest valid completed attempt for each case and trial is used once. No measured outlier is silently discarded.")
    r.p("analysis/coverage.json identifies missing runs and pending adaptive cases. analysis/issues.csv flags invalid data, unusually low queue occupancy estimates, read/write-mix deviations, and small p99.9 sample counts. All submitted plots and tables must be traceable to the retained raw outputs. Host metadata, operator notes, and any unavailable sensors complete the environment record.")

    r.page(); r.h("4 Analysis framework")
    r.h("Metrics and aggregation",2)
    r.p("IOPS counts requests per second. Throughput is retained as MiB/s = bytes/s divided by 2^20 and MB/s = bytes/s divided by 10^6. For a fixed block size, bytes/s is approximately IOPS times bytes/request. Larger blocks can increase bandwidth while reducing request rate, so the two metrics must be interpreted together.")
    r.p("Curves use mean total latency. Baseline and tail percentiles use completion latency, clat; submission latency, slat, is retained separately. Means are weighted by completed-I/O counts across workers and read/write directions. Aggregate percentiles come from merged JSON+ completion-histogram counts. fio bins have finite resolution and are represented by midpoint values [1]. One-second latency logs contain means and are not used to estimate tails.")
    r.p("Tables retain each trial plus its mean, median, sample SD, minimum, and maximum. Plot error bars show mean plus/minus between-trial sample SD when at least two trials exist. SD is neither a confidence interval nor an individual-I/O percentile. Per-trial N is reported with p99.9; N below 100000 gives fewer than 100 expected observations beyond that percentile and receives a diagnostic note.")
    r.h("Queueing and the knee",2)
    r.p("Little's Law relates average work in the system to throughput and residence time: L = lambda W. Using aggregate IOPS and mean completion latency in seconds estimates work within fio's submitted-I/O boundary. The total-latency version also includes submission. For a mixed workload, L is approximately lambda_read W_read + lambda_write W_write. The configured cap is iodepth times numjobs; neither the cap nor this estimate establishes physical NAND concurrency.")
    r.p("The knee heuristic uses median trial throughput for each workload. It selects the earliest adjacent QD pair where both reach at least 90% of peak observed throughput, the next point adds at most 10% throughput, and mean total latency rises at least 25%. It is a candidate for inspection, not proof of a physical-device limit. If no pair qualifies, the report must say the tested range did not establish a knee; tail comparisons then use the highest tested depth. The middle depth is the tested value nearest half the comparison QD.")
    r.h("Mechanisms to evaluate",2)
    r.p("More concurrency can hide latency until a controller, media, interface, CPU, or virtualization limit dominates. Beyond saturation, additional work can increase waiting and tail latency while adding little throughput. Sequential locality, transfer overhead, random mapping work, and shared read/write resources may change the curve. Working-set size can change locality and cache reuse. These are explanations to test against the measurements, not experimental findings supplied in advance.")
    r.p("Flash updates typically write to new locations. Garbage collection can move valid data and erase blocks, competing with foreground requests. Fast SLC write caching can favor short write bursts; a longer write window may expose different behavior as cache space or temperature changes. Correlate any transition with elapsed time, bytes written, and available SSD sensors before assigning a cause.")
    r.p("Any interface-ceiling discussion must use the actual physical SSD bus and link. The processor model and guest virtual SCSI identity do not establish SATA, NVMe, or PCIe generation. The fio depth histogram is retained, but its highest bucket groups 64 and above; it cannot by itself prove achieved QD 128 or 256.")

    r.page(); r.h("5 Baseline and block size results")
    r.p("Measured baseline values are entered below. All latencies are microseconds. The mean column uses total latency; percentile columns use completion latency. Add between-trial SD and trial count from summary.csv when reporting the final values.")
    baseline=[]
    for pattern,block in (("random",4),("sequential",128)):
        for read in (100,0):
            row=r.find("baseline",pattern,read)
            baseline.append([f"{block} KiB {pattern}\n{'read' if read==100 else 'write'}",*(r.number(row,m) for m in ("mean_lat_us","p50_us","p95_us","p99_us","p999_us")), row["n"] if row else "TBD"])
    r.table(["QD 1 workload","Mean total","p50 clat","p95 clat","p99 clat","p99.9 clat","n"],baseline,[2.05,.85,.85,.85,.85,.95,.6],font=8.7)
    r.fig("01_baseline","Figure 1 QD 1 baseline means and completion percentiles for the four workloads")
    r.p("[TBD Compare read and write baselines and central versus tail latency. Explain that QD 1 minimizes application queueing but retains filesystem, virtualization, controller and device overhead.]")
    r.h("Block size and access pattern",2)
    r.fig("02_blocksize","Figure 2 Throughput and mean total latency from the same block-size runs at QD 16")
    r.fig("03_block_iops","Figure 3 IOPS versus block size for random and sequential reads and writes")
    r.p("[TBD Identify the block sizes with the highest observed bandwidth and IOPS, distinguish MiB/s from MB/s, and compare random with sequential patterns. Explain any rise in bandwidth accompanied by lower IOPS using the fixed-size conversion. Discuss repeat variability rather than attributing every small difference to a device effect.]")
    r.p("Maximum observed random 4 KiB IOPS  Read [TBD]  Write [TBD]\nMaximum observed sequential throughput  Read [TBD] MiB/s  Write [TBD] MiB/s\nCase identifiers and fixed knobs for each maximum  [TBD]")
    r.p("Use analysis/maxima.csv to identify the maximum within the tested matrix. This is an observed profile maximum, not a manufacturer specification or proof of the device's absolute capability.")

    r.page(); r.h("6 Mix queue depth and worker results")
    r.fig("04_mix","Figure 4 Read/write mix at fixed QD and block size within each pattern")
    r.p("[TBD Compare 100/0, 70/30, 50/50 and 0/100. Use achieved read fractions and separate read/write rows to explain any asymmetric latency effect. Evaluate possible controller or shared-media contention with evidence.]")
    r.fig("05_queue_tradeoff","Figure 5 Single throughput versus mean total latency curve with QD labels for all four workloads")
    r.fig("06_queue_scaling","Figure 6 Throughput and mean latency as configured QD increases")
    knee_rows=r.csv("knees.csv")
    table=[]
    for pattern in ("random","sequential"):
        for read in (100,0):
            row=next((v for v in knee_rows if v["pattern"]==pattern and int(v["read_pct"])==read),None)
            table.append([pattern+" "+("read" if read==100 else "write"),row["selected_qd"] if row else "TBD",("Candidate" if row["plateau_observed"]=="True" else "Not established") if row else "TBD","TBD"])
    r.table(["Workload","Selected QD","Plateau assessment","Supporting throughput and latency"],table,[1.65,.85,1.4,3.1])
    r.p("[TBD Accept or qualify the heuristic candidate using adjacent points and SD. If no plateau is observed, state that the knee was not established in this range and identify the fallback depth. Explain whether the plateau could reflect host CPU, virtualization, interface or media limits; do not choose a cause without evidence.]")
    r.fig("07_little_law","Figure 7 IOPS times mean completion latency compared with configured aggregate QD")
    r.p("Numeric Little's Law check  Workload and QD [TBD]; lambda [TBD] requests/s; W [TBD] us / 1000000; L [TBD] requests; configured QD [TBD]. Compare submitted and total-latency estimates and discuss underfilling, submission overhead, or nonstationarity. Use analysis/little_law.csv and per-job depth buckets.")
    r.fig("08_parallel","Figure 8 Worker count at fixed aggregate QD and fixed total address range")
    r.p("[TBD Compare one deep queue with several shallower workers at aggregate QD 16 and 64. Relate changes to CPU overhead or submission parallelism while checking achieved occupancy.]")

    r.page(); r.h("7 Tail latency working set and device state")
    tailrows=[]
    for pattern in ("random","sequential"):
        for read in (100,0):
            available=[x for x in r.rows if x["family"]=="tail" and x["direction"]=="total" and x["pattern"]==pattern and int(x["read_pct"])==read]
            for role in ("mid","comparison"):
                row=next((x for x in available if (x["role"]=="mid" if role=="mid" else x["role"]!="mid")),None)
                tailrows.append([f"{pattern}\n{'R' if read==100 else 'W'} {row['role'] if row else role}",row["iodepth"] if row else "TBD",*(r.number(row,m) for m in ("p50_us","p95_us","p99_us","p999_us")),str(int(float(row.get("clat_samples_mean", row["io_count_mean"])))) if row else "TBD"])
    r.p("Completion percentiles below are microseconds and are mean per-trial percentile values. N is the mean completed-I/O count per trial; include its range and SD when interpreting a sparse tail. The comparison depth is near-knee only when the declared rule has evidence for a plateau.")
    r.table(["Workload and role","QD","p50","p95","p99","p99.9","Mean N"],tailrows,[2.05,.55,.8,.8,.8,.9,1.1],font=8.7)
    r.fig("09_tails","Figure 9 Completion-latency percentiles at middle and selected comparison QD")
    r.fig("10_tail_survival","Figure 10 Labeled pooled histogram survival curves across trials")
    r.p("[TBD Compare p99.9/p50 and p99/p50 across depths. State trial counts and any small-N diagnostics. Explain whether increased concurrency mainly changes central latency, tails, or both, and what that means for response-time-sensitive workloads.]")
    r.h("Working set and observed write state",2)
    r.fig("11_working_set","Figure 11 Guest address-range effects at random 4 KiB and QD 16")
    r.fig("12_state","Figure 12 Throughput and mean completion latency during each sustained write window")
    r.p("[TBD Record tested ranges, compare IOPS and p99, and explain that guest address windows do not select physical flash LBAs. For each state trial, compare first and last thirds using state_windows.csv. Pair changes with available SSD temperature or host reliability observations. If no change occurs, limit the conclusion to this duration and range.]")

    r.page(); r.h("8 Misleading setup and limitations")
    r.h("Warm buffered reads as an SSD benchmark pitfall",2)
    r.p("The controlled pitfall experiment uses a finite buffered sequential read to warm the same initialized 256 MiB or smaller range before each case. Both measured cases use random 4 KiB reads, psync, QD 1, one worker, and the same range. Only direct=1 versus direct=0 changes. Holding the engine constant avoids confusing the cache effect with asynchronous submission overhead.")
    r.fig("13_pitfall","Figure 13 Direct and warmed buffered reads under the same synchronous workload")
    r.p("Buffered reads can be satisfied from guest RAM, so presenting their request rate as physical SSD IOPS can be misleading. A cache benefit depends on effective cache residency and the storage layers involved. If the two controls are similar, check whether warming completed, whether the range fits available memory, and whether host/virtualization overhead dominates. No global page-cache flush is performed. direct=1 concerns the guest cache and cannot certify the behavior of every host cache.")
    r.p("[TBD Report each control's IOPS and latency with trial variation, compute buffered/direct IOPS, and explain the observed direction and size of the effect. State whether the experiment demonstrated an effect or only bounded it.]")
    r.h("Storage path and durability",2)
    r.p("WSL2 stores a distribution in an ext4 virtual hard disk [3]. This experiment includes the Linux filesystem, virtual disk and host storage layers. Linux-local storage avoids the separate Windows-mounted file path that Microsoft advises against for Linux performance work [2], but the results still describe a virtualized path. Physical SSD temperature, interface detail and hardware queue depth may not be exposed to the guest. Guest free space may reflect the VHD limit rather than actual remaining host capacity [3].")
    r.p("Write completion is measured without per-I/O fsync or a power-loss test. Consequently these write latencies do not establish durable NAND commit latency. The full-file initialization does not reset physical drive history or force a standardized steady state. Address ranges within a file are mediated by filesystem allocation, VHDX mapping and the flash translation layer.")
    r.h("Attribution and uncertainty",2)
    r.p("SLC caching, garbage collection, thermal throttling, host caching, background I/O and CPU limits can produce similar changes. A throughput drop alone does not distinguish them. Use measured temperature and chronology where available, and label other explanations as hypotheses. The final part of a 120-second write window is an observed late window, not automatically a steady state. Shared drive history also limits independence between repeated trials.")
    r.p("[TBD List unavailable sensors, unexplained anomalies, incomplete settings, interruptions, CPU or memory pressure, and the limitations that materially affect the final conclusion.]")

    r.page(); r.h("9 Conclusions and submission evidence")
    r.p("[TBD Summarize the measured QD 1 latency, maximum observed random IOPS and sequential bandwidth, knee or lack of one, and tail-latency tradeoff. Explain which tested settings best serve latency-sensitive and throughput-oriented work. State the strongest supported device-state observation and the effect of the misleading setup. Keep every numerical claim traceable to a case and acknowledge the WSL scope.]")
    r.h("Rubric coverage",2)
    r.table(["Category","Points","Evidence"],[
        ["Safety and QD 1 baselines",30,"Owned-file guards, initialization evidence, metadata and Section 5"],
        ["Block size and patterns",35,"Complete 24-case matrix; Figures 2 and 3; Section 5"],
        ["Mix and queue analysis",40,"Mix/QD/worker sweeps, knee evidence and numeric Little's Law check; Section 6"],
        ["Tails and device state",25,"p50/p95/p99/p99.9, N, ranges, state chronology and temperature limits; Section 7"],
        ["Reproducibility and hygiene",25,"Configs, raw data, hashes, repeated trials, environment, diagnostics and coverage"],
        ["AI and pitfall analysis",15,"Disclosure below and controlled cache experiment; Section 8"],
        ["Report quality",15,"Labeled figures, units, concise supported interpretation and limitations"],
        ["Total",185,"Final empirical evidence must accompany the implementation"]
    ],[2.0,.55,4.45],font=8.8)
    r.h("AI use disclosure",2)
    r.p("OpenAI ChatGPT assisted in generating the experiment code, parsing and plotting routines, documentation, and report structure from the supplied assignment. The author is responsible for running the target-laptop experiments, verifying raw results and generated statements, and writing the final empirical interpretation. No target-laptop performance results are supplied in the starter report. Update this disclosure to reflect any additional assistance.")
    r.h("References",2)
    refs=[
        "[1] fio project. Official fio documentation. https://fio.readthedocs.io/en/latest/fio_doc.html",
        "[2] Microsoft. Working across Windows and Linux file systems. https://learn.microsoft.com/en-us/windows/wsl/filesystems",
        "[3] Microsoft. How to manage WSL disk space. https://learn.microsoft.com/en-us/windows/wsl/disk-space",
        "[4] Microsoft. Get PhysicalDisk. https://learn.microsoft.com/en-us/powershell/module/storage/get-physicaldisk",
        "[5] Microsoft. Get StorageReliabilityCounter. https://learn.microsoft.com/en-us/powershell/module/storage/get-storagereliabilitycounter",
        "[6] Advanced Computer Systems. Project 3 SSD Performance Profiling assignment and 185-point grading rubric supplied with the project."
    ]
    for text in refs:
        p=r.p(text)
        for run in p.runs: run.font.size=Pt(8)
        p.paragraph_format.space_after=Pt(3)
    output=Path(output)
    output.parent.mkdir(parents=True,exist_ok=True)
    if output.exists() or output.with_suffix(".md").exists():
        raise ValueError("Choose a new output path; existing reports are never overwritten")
    d.save(output)
    output.with_suffix(".md").write_text("\n".join(r.md),encoding="utf-8")
    print(f"Saved {output} and matching Markdown. Empirical interpretation fields remain TBD.")


if __name__=="__main__":
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument("--results",help="Existing results directory containing analysis/ and optional plots/")
    p.add_argument("--output",required=True)
    a=p.parse_args()
    build(a.output,a.results)
