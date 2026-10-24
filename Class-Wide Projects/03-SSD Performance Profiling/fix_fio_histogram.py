#!/usr/bin/env python3
"""Repair the original SSD profiling parser and migrate an existing checkpoint.

Run from ~/ssd_profiling with the project's Python:
  .venv/bin/python fix_fio_histogram.py --results results/full_01

No fio benchmark is executed. Raw results and workload.bin are never written.
The old source files and manifest are backed up, and old/new code provenance
is recorded before resume is enabled. Unexpected source changes are refused.
"""
from __future__ import annotations
import argparse
import hashlib
import json
import subprocess
import sys
from pathlib import Path

OLD_HASHES = {'ssdbench/analysis.py': '80a46277100e1427f519b61edc7504823e2a5800d6ccbb8f7fd50ebad61beb5d', 'scripts/build_report.py': 'bdc4f676cbbf30f12712771b87f24d06a7e8edc76163d0d8af92278a5ce32b23'}
NEW_HASHES = {'ssdbench/analysis.py': 'e49de90e08f6c4e8a2b248e84a66812c1dbda56b0f65696b45207380157d254e', 'scripts/build_report.py': '31f06b532a0e79aa770fca5f1cdf3095919ed1f027fd530b81e0dec42e147a1c'}
PAYLOADS = {}
PAYLOADS['ssdbench/analysis.py'] = r'''from __future__ import annotations

import csv
import hashlib
import json
import math
import statistics as stats
from collections import defaultdict
from pathlib import Path

from .common import dump_json, read_json, write_csv

METRICS = ("iops", "mib_s", "mb_s", "mean_lat_us", "mean_clat_us", "mean_slat_us",
           "p50_us", "p95_us", "p99_us", "p999_us", "io_count", "clat_samples", "io_bytes",
           "achieved_read_pct", "little_l_total", "little_l_clat", "clat_qd_fraction")
META = ("case_id", "family", "pattern", "read_pct", "bs", "iodepth", "numjobs", "total_qd",
        "working_set", "runtime", "ramp", "direct", "engine", "role", "direction")


def latency_block(stat, kind):
    for unit, factor in (("ns", .001), ("us", 1.), ("ms", 1000.)):
        if f"{kind}_{unit}" in stat:
            return stat[f"{kind}_{unit}"], factor
    return {}, 1.


def histogram_percentile(bins, percentile):
    n = sum(bins.values())
    if n == 0:
        return math.nan
    threshold, cumulative = math.ceil(n * percentile / 100), 0
    for latency, count in sorted(bins.items()):
        cumulative += count
        if cumulative >= threshold:
            return latency
    return max(bins)


def parse_document(doc, case, run_id, trial):
    jobs = doc.get("jobs", [])
    if len(jobs) != case["numjobs"]:
        raise ValueError("Expected per-job JSON (group_reporting=0); job count differs from numjobs")
    if any(int(j.get("error", 0)) for j in jobs):
        raise ValueError("fio reported a nonzero job error")
    direction_data = defaultdict(list)
    issues = []
    for job_index, job in enumerate(jobs, 1):
        boundary_excess = 0
        for direction in ("read", "write"):
            st = job.get(direction, {})
            n = int(st.get("total_ios", 0))
            if n:
                block, _ = latency_block(st, "clat")
                samples = int(block.get("N", n))
                bins = block.get("bins", {})
                if bins and sum(int(v) for v in bins.values()) != samples:
                    raise ValueError("Completion histogram count differs from completion-latency sample N")
                delta = samples - n
                if delta:
                    if not (case["ramp"] > 0 and case["engine"] == "libaio"
                            and 0 < delta <= max(0, case["iodepth"] - 1)):
                        raise ValueError("Latency sample count differs from I/O count outside the bounded asynchronous ramp case")
                    boundary_excess += delta
                    issues.append({"run_id": run_id, "type": "ramp_boundary_accounting",
                                   "detail": f"Job {job_index} {direction}: total_ios={n}, clat_N={samples}, excess={delta}. Histogram agrees with clat_N; bounded in-flight ramp-boundary samples retained."})
                direction_data[direction].append(st)
        if boundary_excess > max(0, case["iodepth"] - 1):
            raise ValueError("Combined read/write latency excess exceeds the per-job ramp-boundary bound")
    if not direction_data:
        raise ValueError("No completed I/O")
    if case["read_pct"] == 100 and "write" in direction_data or case["read_pct"] == 0 and "read" in direction_data:
        raise ValueError("Observed I/O direction disagrees with the configured pure workload")
    if 0 < case["read_pct"] < 100 and len(direction_data) != 2:
        raise ValueError("Mixed workload has no completed I/O in one direction")
    read_n = sum(x["total_ios"] for x in direction_data.get("read", []))
    all_n = sum(x["total_ios"] for entries in direction_data.values() for x in entries)
    read_pct = 100 * read_n / all_n
    rows, hist_rows, depth_rows = [], [], []
    direction_data["total"] = [x for d in ("read", "write") for x in direction_data.get(d, [])]
    for direction, entries in direction_data.items():
        count = sum(int(x["total_ios"]) for x in entries)
        def mean(kind):
            observations = [(latency_block(x, kind), x["total_ios"]) for x in entries]
            denominator = sum(int(block.get("N", fallback)) for (block, _), fallback in observations)
            return (sum(block.get("mean", 0.) * factor * int(block.get("N", fallback))
                        for (block, factor), fallback in observations) / denominator) if denominator else 0.
        clat_samples = sum(int(latency_block(x, "clat")[0].get("N", x["total_ios"])) for x in entries)
        bins = defaultdict(int)
        for x in entries:
            block, factor = latency_block(x, "clat")
            for value, n in block.get("bins", {}).items():
                bins[float(value) * factor] += int(n)
        if len(entries) > 1 and not bins:
            raise ValueError("Cannot merge latency percentiles without JSON+ histogram bins")
        if bins and sum(bins.values()) != clat_samples:
            raise ValueError("Merged completion histogram count differs from completion-latency sample N")
        pct = {}
        for p, key in ((50, "p50_us"), (95, "p95_us"), (99, "p99_us"), (99.9, "p999_us")):
            if bins:
                pct[key] = histogram_percentile(bins, p)
            else:
                block, factor = latency_block(entries[0], "clat")
                vals = block.get("percentile", {})
                pct[key] = next((float(v)*factor for k, v in vals.items() if abs(float(k)-p)<1e-5), math.nan)
        iops = sum(float(x.get("iops", 0.)) for x in entries)
        bw = sum(float(x.get("bw_bytes", x.get("bw", 0.) * 1024)) for x in entries)
        qd = case["iodepth"] * case["numjobs"]
        row = {**case, "direction": direction, "run_id": run_id, "trial": trial,
               "io_count": count, "clat_samples": clat_samples, "io_bytes": sum(int(x.get("io_bytes", x.get("io_kbytes", 0)*1024)) for x in entries),
               "iops": iops, "mib_s": bw / 1024**2, "mb_s": bw / 1e6,
               "mean_lat_us": mean("lat"), "mean_clat_us": mean("clat"), "mean_slat_us": mean("slat"),
               "achieved_read_pct": read_pct, **pct}
        row["little_l_total"] = iops * row["mean_lat_us"] / 1e6
        row["little_l_clat"] = iops * row["mean_clat_us"] / 1e6
        row["clat_qd_fraction"] = row["little_l_clat"] / qd
        rows.append(row)
        for latency, n in sorted(bins.items()):
            hist_rows.append({"run_id": run_id, "case_id": case["case_id"], "trial": trial,
                              "direction": direction, "clat_us": latency, "count": n})
        if direction != "total" and clat_samples < 100000:
            issues.append({"run_id": run_id, "type": "tail_sample_size", "detail": f"{direction}: clat N={clat_samples}; fewer than 100 expected samples beyond p99.9. Interpret that tail cautiously."})
    total = next(x for x in rows if x["direction"] == "total")
    if case["engine"] == "libaio" and qd > 1 and total["clat_qd_fraction"] < .5:
        issues.append({"run_id": run_id, "type": "queue_occupancy", "detail": f"lambda*mean(clat)={total['little_l_clat']:.2f} versus configured aggregate QD={qd}; inspect per-job depth buckets, submission overhead, and virtualization."})
    if abs(read_pct - case["read_pct"]) > 5 and all_n > 10000:
        issues.append({"run_id": run_id, "type": "mix_deviation", "detail": f"Achieved read percentage {read_pct:.2f}, requested {case['read_pct']}"})
    for index, job in enumerate(jobs, 1):
        for bucket, pct in job.get("iodepth_level", {}).items():
            depth_rows.append({"run_id": run_id, "case_id": case["case_id"], "trial": trial,
                               "job": index, "configured_per_job_qd": case["iodepth"],
                               "bucket": str(bucket), "percent": pct,
                               "note": "64 means at least 64; buckets are not exact time-weighted hardware occupancy"})
    return rows, hist_rows, depth_rows, issues


def summarize_rows(rows):
    groups = defaultdict(list)
    for row in rows:
        groups[(row["case_id"], row["direction"])].append(row)
    summaries = []
    for key, group in sorted(groups.items()):
        result = {k: group[0].get(k, "") for k in META}
        result["n"] = len(group)
        for metric in METRICS:
            values = [float(r[metric]) for r in group if math.isfinite(float(r[metric]))]
            for suffix, value in (("mean", stats.mean(values) if values else math.nan),
                                  ("median", stats.median(values) if values else math.nan),
                                  ("sd", stats.stdev(values) if len(values)>1 else 0.),
                                  ("min", min(values) if values else math.nan),
                                  ("max", max(values) if values else math.nan)):
                result[f"{metric}_{suffix}"] = value
        summaries.append(result)
    return summaries


def choose_knee(curve):
    """A declared saturation heuristic, never an assertion of a physical device knee."""
    curve = sorted(curve, key=lambda x: x["iodepth"])
    maximum = max(r["mib_s_median"] for r in curve)
    for a, b in zip(curve, curve[1:]):
        if (maximum > 0 and a["mib_s_median"] >= .9*maximum and b["mib_s_median"] >= .9*maximum
                and b["mib_s_median"] <= 1.1*a["mib_s_median"]
                and b["mean_lat_us_median"] >= 1.25*a["mean_lat_us_median"]):
            return {"selected_qd": a["iodepth"], "plateau_observed": True,
                    "reason": "Earliest adjacent pair at >=90% of peak median throughput, <=10% throughput gain, and >=25% latency growth"}
    return {"selected_qd": curve[-1]["iodepth"], "plateau_observed": False,
            "reason": "No adjacent pair satisfied the declared plateau rule; use highest tested QD for tail comparison and report knee not established"}


def load_events(out):
    path = out / "journal.jsonl"
    if not path.exists():
        return []
    events = []
    lines = path.read_text().splitlines()
    for i, line in enumerate(lines):
        try:
            events.append(json.loads(line))
        except ValueError:
            if i != len(lines)-1:
                # Runner repairs a truncated record by adding a newline; retain an explicit issue.
                events.append({"status": "corrupt", "error": f"Invalid journal line {i+1}"})
    return events


def time_logs(out, event):
    rows = []
    for kind, metric, factor in (("bw", "mib_s", 1/1024), ("iops", "iops", 1),
                                 ("clat", "mean_clat_us", 1/1000), ("lat", "mean_lat_us", 1/1000)):
        for path in sorted((out / "logs").glob(f"{event['run_id']}_{kind}.*.log")):
            job = path.name.rsplit(".", 2)[-2]
            with path.open() as f:
                for line in csv.reader(f):
                    if len(line) < 3:
                        continue
                    try:
                        rows.append({"run_id": event["run_id"], "case_id": event["case"]["case_id"],
                                     "family": event["case"]["family"], "trial": event["trial"], "job": job,
                                     "time_s": float(line[0])/1000,
                                     "direction": {0:"read", 1:"write", 2:"trim"}.get(int(line[2]), "other"),
                                     "metric": metric, "value": float(line[1])*factor})
                    except (ValueError, IndexError):
                        raise ValueError(f"Malformed fio time log: {path}")
    return rows


def markdown_table(path, headers, rows):
    lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"]
    lines += ["| " + " | ".join(str(v) for v in row) + " |" for row in rows]
    path.write_text("\n".join(lines) + "\n")


def interval_window_mean(values, start, end):
    """Weight interval means by overlap with a fixed observation horizon.

    fio may emit a tiny final interval during queue drain. It remains in the
    raw logs and plots, but cannot receive the weight of a full second here.
    """
    previous = 0.
    weighted = duration = 0.
    for row in sorted(values,key=lambda r:r["time_s"]):
        timestamp = row["time_s"]
        overlap = max(0.,min(timestamp,end)-max(previous,start))
        weighted += row["value"] * overlap
        duration += overlap
        previous = max(previous,timestamp)
    return weighted/duration if duration else ""


def analyze(directory):
    out = Path(directory)
    dest = out / "analysis"
    dest.mkdir(exist_ok=True)
    events = load_events(out)
    completed = {}
    for e in events:
        if e.get("phase") == "measured" and e.get("status") == "complete":
            completed[(e["case"]["case_id"], e["trial"])] = e
    rows, hist, depths, issues, times = [], [], [], [], []
    valid = set()
    for key, e in sorted(completed.items()):
        try:
            raw = out / e["raw"]
            if e.get("raw_sha256") and hashlib.sha256(raw.read_bytes()).hexdigest() != e["raw_sha256"]:
                raise ValueError("Raw SHA256 differs from journal")
            r, h, d, i = parse_document(read_json(raw), e["case"], e["run_id"], e["trial"])
            t = time_logs(out, e)
            rows.extend(r); hist.extend(h); depths.extend(d); issues.extend(i); times.extend(t)
            valid.add(key)
        except (OSError, ValueError, KeyError, TypeError) as exc:
            issues.append({"run_id": e["run_id"], "type": "invalid_result", "detail": str(exc)})
    for e in events:
        if e.get("status") in ("failed", "corrupt"):
            issues.append({"run_id": e.get("run_id", "journal"), "type": "execution_failure", "detail": e.get("error", "unknown")})
    summary = summarize_rows(rows)
    write_csv(dest / "measurements.csv", rows, [*META, "run_id", "trial", *METRICS])
    write_csv(dest / "summary.csv", summary, [*META, "n", *(f"{m}_{s}" for m in METRICS for s in ("mean", "median", "sd", "min", "max"))])
    write_csv(dest / "histograms.csv", hist, ["run_id", "case_id", "trial", "direction", "clat_us", "count"])
    write_csv(dest / "depths.csv", depths, ["run_id", "case_id", "trial", "job", "configured_per_job_qd", "bucket", "percent", "note"])
    write_csv(dest / "issues.csv", issues, ["run_id", "type", "detail"])
    write_csv(dest / "timeseries.csv", times, ["run_id", "case_id", "family", "trial", "job", "time_s", "direction", "metric", "value"])
    knees = []
    for pattern in ("random", "sequential"):
        for read in (100, 0):
            curve = [r for r in summary if r["family"]=="queue" and r["direction"]=="total" and r["pattern"]==pattern and r["read_pct"]==read]
            if curve:
                knees.append({"pattern": pattern, "read_pct": read, **choose_knee(curve)})
    write_csv(dest / "knees.csv", knees, ["pattern", "read_pct", "selected_qd", "plateau_observed", "reason"])
    little = [{k:r[k] for k in ("case_id", "pattern", "read_pct", "iodepth", "numjobs", "total_qd",
                "iops_mean", "mean_lat_us_mean", "mean_clat_us_mean", "little_l_total_mean", "little_l_clat_mean", "clat_qd_fraction_mean")}
              for r in summary if r["direction"]=="total" and r["family"] in ("queue", "parallel")]
    write_csv(dest / "little_law.csv", little)
    state = []
    for key, e in sorted(completed.items()):
        if key not in valid or e["case"]["family"] != "state":
            continue
        for metric in ("mib_s", "mean_clat_us"):
            values = sorted([t for t in times if t["run_id"]==e["run_id"] and t["metric"]==metric and t["direction"]=="write"], key=lambda t:t["time_s"])
            if not values:
                continue
            # Equal-duration thirds (not an assertion that the last third is a steady state).
            duration = e["case"]["runtime"]
            state.append({"run_id":e["run_id"], "trial":e["trial"], "metric":metric,
                          "first_third_mean":interval_window_mean(values,0,duration/3),
                          "last_third_mean":interval_window_mean(values,2*duration/3,duration),
                          "note":"Interval means weighted by time overlap with measured horizon; queue drain after runtime has zero weight. No steady-state claim."})
    write_csv(dest / "state_windows.csv", state, ["run_id", "trial", "metric", "first_third_mean", "last_third_mean", "note"])
    manifest = read_json(out / "manifest.json") if (out / "manifest.json").exists() else {"parameters":{}}
    params = manifest["parameters"]
    expected_cases = read_json(out / "plan.json")["cases"] if (out / "plan.json").exists() else []
    if (out / "adaptive_tails.json").exists():
        expected_cases += read_json(out / "adaptive_tails.json")["cases"]
    expected = {(c["case_id"],t) for c in expected_cases for t in range(1, params.get("trials",1)+1)}
    coverage = {"expected_runs_selected_so_far":len(expected), "valid_runs":len(valid),
                "missing_runs":[{"case_id":k[0],"trial":k[1]} for k in sorted(expected-valid)],
                "adaptive_tails_pending":"tail" in params.get("families",[]) and not (out/"adaptive_tails.json").exists(),
                "requested_families":params.get("families",[]), "profile":params.get("profile", "unknown"),
                "all_selected_runs_complete":expected.issubset(valid) and bool(expected),
                "warnings":len(issues), "measured_write_gib_excluding_ramp_and_initialization":sum(r["io_bytes"] for r in rows if r["direction"]=="write")/1024**3}
    coverage["all_selected_runs_complete"] &= not coverage["adaptive_tails_pending"]
    # A successful subset or smoke run must never imply full assignment coverage.
    from .plan import build_cases
    required_ids = {c.case_id for c in build_cases(params.get("size",16*1024**3))}
    repeated_ids = {cid for cid in {k[0] for k in valid} if sum(k[0]==cid for k in valid)>=2}
    tail_ids = {c["case_id"] for c in expected_cases if c["family"]=="tail"}
    coverage["full_default_setting_matrix_with_repeats"] = required_ids.issubset(repeated_ids) and len(tail_ids)==8 and tail_ids.issubset(repeated_ids)
    coverage["coverage_note"] = "Selected-run completion and matrix coverage are separate checks; neither establishes report quality or causal interpretation."
    dump_json(dest/"coverage.json", coverage)
    baselines = [r for r in summary if r["family"]=="baseline" and r["direction"]!="total"]
    markdown_table(dest/"baseline_table.md", ["Pattern", "KiB", "Direction", "n", "IOPS", "MiB/s", "Mean total us", "p50 clat us", "p95 clat us", "p99 clat us", "p99.9 clat us"],
                   [[r["pattern"], r["bs"]//1024, r["direction"], r["n"], *(f"{r[m+'_mean']:.3f}" for m in ("iops","mib_s","mean_lat_us","p50_us","p95_us","p99_us","p999_us"))] for r in baselines])
    tails = [r for r in summary if r["family"]=="tail" and r["direction"]!="total"]
    markdown_table(dest/"tail_table.md", ["Pattern", "Direction", "Role", "QD", "n", "Mean N", "p50 us", "p95 us", "p99 us", "p99.9 us"],
                   [[r["pattern"],r["direction"],r["role"],r["iodepth"],r["n"],int(r["clat_samples_mean"]),*(f"{r[m+'_mean']:.3f}" for m in ("p50_us","p95_us","p99_us","p999_us"))] for r in tails])
    maxima=[]
    for read in (100,0):
        for name, pattern, bs, metric in (("small_block_iops","random",4096,"iops"),("sequential_throughput","sequential",None,"mib_s")):
            candidates=[r for r in summary if r["direction"]=="total" and r["read_pct"]==read and r["pattern"]==pattern and (bs is None or r["bs"]==bs) and r["direct"]==1 and r["family"] in ("block","queue","parallel")]
            if candidates:
                r=max(candidates,key=lambda r:r[metric+"_median"])
                maxima.append({"measurement":name,"read_pct":read,"metric":metric,"median":r[metric+"_median"],"case_id":r["case_id"],"note":"Maximum observed in tested matrix, not a device specification"})
    write_csv(dest/"maxima.csv",maxima,["measurement","read_pct","metric","median","case_id","note"])
    print(f"Parsed {len(valid)} valid measured runs; {len(expected-valid)} missing; {len(issues)} diagnostic notes. {dest}")
    return coverage
'''
PAYLOADS['scripts/build_report.py'] = r'''#!/usr/bin/env python3
"""Create an editable report, optionally inserting measured tables and figures.

Interpretation remains the student's work. Existing documents are never overwritten.
"""
from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


class Report:
    def __init__(self, results=None):
        self.doc=Document()
        # The bundled default template may carry decorative title borders.
        for element in self.doc.styles.element.xpath(".//w:pBdr"):
            element.getparent().remove(element)
        self.md=[]
        self.results=Path(results) if results else None
        self.rows=self.csv("summary.csv")
        section=self.doc.sections[0]
        section.page_width=Inches(8.5); section.page_height=Inches(11)
        section.top_margin=section.bottom_margin=Inches(.7)
        section.left_margin=section.right_margin=Inches(.75)
        section.header_distance=section.footer_distance=Inches(.28)
        for name,size in (("Normal",10.5),("Title",24),("Subtitle",11),("Heading 1",14),("Heading 2",11.5),("Caption",9)):
            st=self.doc.styles[name]
            st.font.name="Arial"; st.font.size=Pt(size); st.font.color.rgb=RGBColor(0,0,0)
            st.paragraph_format.space_after=Pt(5)
            st.paragraph_format.line_spacing=1.08
            if name.startswith("Heading"):
                st.paragraph_format.space_before=Pt(10)
                st.paragraph_format.keep_with_next=True
            if name=="Title": st.paragraph_format.space_after=Pt(8)
        header=section.header.paragraphs[0]
        header.text="SSD Performance Profiling"
        header.runs[0].font.size=Pt(8); header.runs[0].font.color.rgb=RGBColor(0,0,0)
        footer=section.footer.paragraphs[0]
        footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run("Page ").font.size=Pt(8)
        field=OxmlElement("w:fldSimple"); field.set(qn("w:instr"),"PAGE"); footer._p.append(field)
        self.doc.core_properties.title="SSD Performance Profiling"
        self.doc.core_properties.subject="Advanced Computer Systems Project 3"
        self.doc.core_properties.author=""
        self.doc.core_properties.keywords="SSD fio WSL performance queue depth tail latency"

    def csv(self,name):
        p=self.results/"analysis"/name if self.results else None
        if not p or not p.exists(): return []
        with p.open(newline="") as f: return list(csv.DictReader(f))

    def p(self,text,italic=False):
        paragraph=self.doc.add_paragraph()
        run=paragraph.add_run(text); run.italic=italic
        if italic: run.font.size=Pt(9.5)
        self.md.extend([text,""])
        return paragraph

    def h(self,text,level=1):
        self.doc.add_heading(text,level)
        self.md.extend(["#"*(level+1)+" "+text,""])

    def page(self): self.doc.add_page_break()

    def table(self,headers,rows,widths=None,font=9):
        table=self.doc.add_table(rows=1,cols=len(headers))
        table.alignment=WD_TABLE_ALIGNMENT.CENTER
        table.autofit=False
        widths=widths or [7/len(headers)]*len(headers)
        for col,width in zip(table.columns,widths): col.width=Inches(width)
        for cell,text,width in zip(table.rows[0].cells,headers,widths):
            cell.text=str(text); cell.width=Inches(width)
            shade=OxmlElement("w:shd"); shade.set(qn("w:fill"),"DCE6F1"); cell._tc.get_or_add_tcPr().append(shade)
            for run in cell.paragraphs[0].runs: run.bold=True
        repeat=OxmlElement("w:tblHeader"); table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            cells=table.add_row().cells
            for cell,text,width in zip(cells,row,widths):
                cell.text=str(text); cell.width=Inches(width)
        for row_index,row in enumerate(table.rows):
            no_split=OxmlElement("w:cantSplit"); row._tr.get_or_add_trPr().append(no_split)
            for j,cell in enumerate(row.cells):
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr=cell._tc.get_or_add_tcPr()
                borders=OxmlElement("w:tcBorders")
                for edge in ("top","left","bottom","right"):
                    e=OxmlElement("w:"+edge); e.set(qn("w:val"),"single"); e.set(qn("w:sz"),"4"); e.set(qn("w:color"),"D9D9D9"); borders.append(e)
                pr.append(borders)
                margins=OxmlElement("w:tcMar")
                for edge in ("top","left","bottom","right"):
                    e=OxmlElement("w:"+edge); e.set(qn("w:w"),"85"); e.set(qn("w:type"),"dxa"); margins.append(e)
                pr.append(margins)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(1); p.paragraph_format.space_before=Pt(1)
                    p.paragraph_format.line_spacing=1.05
                    text=cell.text
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or (len(text)>28 and row_index>0) else WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs: run.font.size=Pt(font); run.font.color.rgb=RGBColor(0,0,0)
        self.md.append("| "+" | ".join(headers)+" |")
        self.md.append("| "+" | ".join("---" for _ in headers)+" |")
        self.md.extend("| "+" | ".join(str(v).replace("\n"," ") for v in row)+" |" for row in rows)
        self.md.append("")
        self.doc.add_paragraph().paragraph_format.space_after=Pt(0)

    def fig(self,stem,caption):
        path=self.results/"plots"/(stem+".png") if self.results else None
        if path and path.exists():
            p=self.doc.add_paragraph(); p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(path),width=Inches(6.7))
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            self.doc.add_paragraph(caption,style="Caption")
            self.md.extend([f"![{caption}]({path.absolute()})",""])
        else:
            self.p("[Insert "+stem+".png] "+caption,italic=True)

    def find(self,family,pattern,read,role=None):
        return next((r for r in self.rows if r["family"]==family and r["direction"]=="total" and r["pattern"]==pattern and int(r["read_pct"])==read and (role is None or r["role"]==role)),None)

    @staticmethod
    def number(row,key,sd=False):
        if not row: return "TBD"
        if sd and int(row["n"])>1: return f"{float(row[key+'_mean']):.2f} +/- {float(row[key+'_sd']):.2f}"
        return f"{float(row[key+'_mean']):.2f}"


def build(output,results=None):
    r=Report(results)
    d=r.doc
    d.add_paragraph("SSD Performance Profiling",style="Title")
    d.add_paragraph("Advanced Computer Systems Project 3",style="Subtitle")
    r.md.extend(["# SSD Performance Profiling","","Advanced Computer Systems Project 3",""])
    r.p("Author  [TBD]     Course section  [TBD]     Measurement date  [TBD]")
    r.h("1 Objective and scope")
    r.p("This project characterizes how block size, access pattern, read/write mix, queue depth, worker count, and address range affect storage throughput and latency. It measures QD 1 baselines, maximum observed small-block IOPS and sequential bandwidth, the throughput versus latency tradeoff, and completion-latency tails. The experiment also tests a warm-page-cache pitfall and observes performance during sustained writes.")
    r.p("The target is a Windows 11 laptop using WSL2. The measurements characterize its complete virtualized storage path. The physical SSD model and interface must be established from host metadata before drawing device-specific conclusions.")
    r.h("System and storage metadata",2)
    r.table(["Item","Configuration"],[
        ["Processor","AMD Ryzen 5 PRO 7530U with Radeon Graphics; advertised base 2.00 GHz"],
        ["Host operating system","Windows 11 x64; build [TBD]"],
        ["WSL environment","WSL version, Linux distribution, kernel, visible CPUs, memory and swap [TBD]"],
        ["Physical SSD","Model, firmware, capacity, bus and link details [TBD]"],
        ["Test path and filesystem","[TBD]; planned owned file under ~/ssdbench-data on Linux ext4"],
        ["Storage mapping","Windows volume hosting the distribution VHDX and its physical SSD [TBD]"],
        ["Capacity and state","Test range [TBD]; host and guest free capacity before/after [TBD]"],
        ["Power and temperature","AC power, Windows power mode, cooling, SSD temperature or unavailable [TBD]"]
    ],[1.55,5.45],font=9)
    r.h("Safety and initialization",2)
    r.p("The runner creates only its own disposable regular file in a new or empty directory. It refuses device paths, symbolic links, hard links, unmanaged files, and changed file identities. One full sequential direct-I/O write with a final flush initializes the tested range before reads. No raw partition, trim, or secure-erase operation is used. This is file preparation, not full-drive preconditioning.")

    r.page(); r.h("2 Experimental design")
    r.p("The full profile uses the matrix below. The planned total file size is 16 GiB. Each ordinary run has a 3-second ramp and 15 measured seconds; each case is repeated three times. Tail runs have 45 measured seconds, and the state observation lasts 120 seconds with no ramp. The actual values and any overrides are retained in manifest.json.")
    r.table(["Experiment","Varied setting","Fixed settings","Cases"],[
        ["QD 1 baseline","Read/write; random/sequential","4 KiB random or 128 KiB sequential; one worker",4],
        ["Block size","4, 16, 32, 64, 128, 256 KiB; both patterns; read/write","QD 16; one worker",24],
        ["Read/write mix","100/0, 70/30, 50/50, 0/100","4 KiB random or 128 KiB sequential; QD 16",8],
        ["Queue depth","1, 2, 4, 8, 16, 32, 64, 128, 256","Four baseline workloads; one worker",36],
        ["Workers","1, 2, 4, 8 workers","Aggregate QD 16 or 64; random 4 KiB read/write",16],
        ["Address range","1/16, 1/4, full file","Random 4 KiB read/write; QD 16",6],
        ["Cache pitfall","Direct or buffered","Warmed range; random 4 KiB reads; psync; QD 1",2],
        ["State observation","Elapsed time in write window","Sequential 128 KiB write; QD 32",1],
        ["Adaptive tails","Middle and near-knee QD","Four baseline workloads; one worker",8]
    ],[1.15,2.25,3.1,.5],font=8.8)
    r.p("The matrix has 97 static cases and eight adaptive cases, giving 315 measured runs at three trials. The default schedule takes about 122 minutes plus initialization, process startup, and cache warming. It is write intensive; the plan displays write exposure before execution. A smoke profile checks operation with fewer settings and does not establish full rubric coverage.")
    r.h("Controls and execution order",2)
    r.p("Main experiments use libaio, direct=1, aligned fixed-size transfers, and fio-generated incompressible buffers. Workers use disjoint regions while holding the total address range constant. The pitfall uses psync for both controls to isolate the direct-I/O setting. Payload refill can consume CPU and is part of this workload. Writes measure completion without per-operation fsync.")
    r.p("Each trial visits the experiment families in the table order, with deterministically shuffled cases within each static family. Tail depths are selected after all queue trials. Trial-specific random seeds and actual execution order are recorded. A two-second pause separates measured runs. Shared drive history, prior writes, and incomplete thermal recovery remain possible sources of dependence.")

    r.page(); r.h("3 Code structure and reproducibility")
    r.table(["Component","Responsibility"],[
        ["scripts/ssdbench.py","Entry point for inspect, init, plan, run, analyze, plot, and cleanup"],
        ["ssdbench/plan.py","Case matrix, alignment, disjoint worker ranges, and fio configuration generation"],
        ["ssdbench/safety.py","Ownership marker, file identity checks, locking, and guarded deletion"],
        ["ssdbench/runner.py","Initialization, trials, exact commands, failures, checkpoints, resume, and adaptive tails"],
        ["ssdbench/analysis.py","Unit conversion, directional metrics, merged histograms, trial statistics, knee rule and coverage"],
        ["ssdbench/plotting.py","Labeled figures with units and repeated-trial SD bars where applicable"],
        ["ssdbench/environment.py","Guest kernel, CPU, mount, memory, disk-counter and available temperature records"],
        ["scripts/collect_host.ps1","Read-only Windows disk, power, free-space and optional reliability metadata"],
        ["scripts/build_report.py","Editable report generation and optional insertion of measured tables and figures"],
        ["tests/test_project.py","Safety, matrix completeness, aggregation, quantiles, statistics and knee selection checks"]
    ],[2.05,4.95],font=9)
    r.h("Workflow",2)
    r.p("Inspect the WSL mount and host storage mapping, fill operator notes, and review the generated plan. Initialize the disposable file once. Check the smoke profile, then run the full profile once; that command already includes repeated trials. Preserve the output directory, regenerate analysis and figures as needed, and add the empirical interpretation to this report. Resume uses the same command plus --resume and verifies settings, source hashes, fio version, target identity, and retained raw hashes. The Windows collector uses the disk and reliability interfaces documented by Microsoft [4, 5].")
    r.h("Data hygiene",2)
    r.p("Each attempt retains its exact fio configuration, unmodified JSON+ output, stderr, auxiliary stdout, command, seed, elapsed time, and SHA256 hash. The journal records actual order and status. Before/after snapshots retain available guest temperatures, memory and disk counters. Failed attempts stay archived and are excluded from analysis. The latest valid completed attempt for each case and trial is used once. No measured outlier is silently discarded.")
    r.p("analysis/coverage.json identifies missing runs and pending adaptive cases. analysis/issues.csv flags invalid data, unusually low queue occupancy estimates, read/write-mix deviations, and small p99.9 sample counts. All submitted plots and tables must be traceable to the retained raw outputs. Host metadata, operator notes, and any unavailable sensors complete the environment record.")

    r.page(); r.h("4 Analysis framework")
    r.h("Metrics and aggregation",2)
    r.p("IOPS counts requests per second. Throughput is retained as MiB/s = bytes/s divided by 2^20 and MB/s = bytes/s divided by 10^6. For a fixed block size, bytes/s is approximately IOPS times bytes/request. Larger blocks can increase bandwidth while reducing request rate, so the two metrics must be interpreted together.")
    r.p("Curves use mean total latency. Baseline and tail percentiles use completion latency, clat; submission latency, slat, is retained separately. Means are weighted by completed-I/O counts across workers and read/write directions. Aggregate percentiles come from merged JSON+ completion-histogram counts. fio bins have finite resolution and are represented by midpoint values [1]. One-second latency logs contain means and are not used to estimate tails.")
    r.p("Tables retain each trial plus its mean, median, sample SD, minimum, and maximum. Plot error bars show mean plus/minus between-trial sample SD when at least two trials exist. SD is neither a confidence interval nor an individual-I/O percentile. Per-trial N is reported with p99.9; N below 100000 gives fewer than 100 expected observations beyond that percentile and receives a diagnostic note.")
    r.h("Queueing and the knee",2)
    r.p("Little's Law relates average work in the system to throughput and residence time: L = lambda W. Using aggregate IOPS and mean completion latency in seconds estimates work within fio's submitted-I/O boundary. The total-latency version also includes submission. For a mixed workload, L is approximately lambda_read W_read + lambda_write W_write. The configured cap is iodepth times numjobs; neither the cap nor this estimate establishes physical NAND concurrency.")
    r.p("The knee heuristic uses median trial throughput for each workload. It selects the earliest adjacent QD pair where both reach at least 90% of peak observed throughput, the next point adds at most 10% throughput, and mean total latency rises at least 25%. It is a candidate for inspection, not proof of a physical-device limit. If no pair qualifies, the report must say the tested range did not establish a knee; tail comparisons then use the highest tested depth. The middle depth is the tested value nearest half the comparison QD.")
    r.h("Mechanisms to evaluate",2)
    r.p("More concurrency can hide latency until a controller, media, interface, CPU, or virtualization limit dominates. Beyond saturation, additional work can increase waiting and tail latency while adding little throughput. Sequential locality, transfer overhead, random mapping work, and shared read/write resources may change the curve. Working-set size can change locality and cache reuse. These are explanations to test against the measurements, not experimental findings supplied in advance.")
    r.p("Flash updates typically write to new locations. Garbage collection can move valid data and erase blocks, competing with foreground requests. Fast SLC write caching can favor short write bursts; a longer write window may expose different behavior as cache space or temperature changes. Correlate any transition with elapsed time, bytes written, and available SSD sensors before assigning a cause.")
    r.p("Any interface-ceiling discussion must use the actual physical SSD bus and link. The processor model and guest virtual SCSI identity do not establish SATA, NVMe, or PCIe generation. The fio depth histogram is retained, but its highest bucket groups 64 and above; it cannot by itself prove achieved QD 128 or 256.")

    r.page(); r.h("5 Baseline and block size results")
    r.p("Measured baseline values are entered below. All latencies are microseconds. The mean column uses total latency; percentile columns use completion latency. Add between-trial SD and trial count from summary.csv when reporting the final values.")
    baseline=[]
    for pattern,block in (("random",4),("sequential",128)):
        for read in (100,0):
            row=r.find("baseline",pattern,read)
            baseline.append([f"{block} KiB {pattern}\n{'read' if read==100 else 'write'}",*(r.number(row,m) for m in ("mean_lat_us","p50_us","p95_us","p99_us","p999_us")), row["n"] if row else "TBD"])
    r.table(["QD 1 workload","Mean total","p50 clat","p95 clat","p99 clat","p99.9 clat","n"],baseline,[2.05,.85,.85,.85,.85,.95,.6],font=8.7)
    r.fig("01_baseline","Figure 1 QD 1 baseline means and completion percentiles for the four workloads")
    r.p("[TBD Compare read and write baselines and central versus tail latency. Explain that QD 1 minimizes application queueing but retains filesystem, virtualization, controller and device overhead.]")
    r.h("Block size and access pattern",2)
    r.fig("02_blocksize","Figure 2 Throughput and mean total latency from the same block-size runs at QD 16")
    r.fig("03_block_iops","Figure 3 IOPS versus block size for random and sequential reads and writes")
    r.p("[TBD Identify the block sizes with the highest observed bandwidth and IOPS, distinguish MiB/s from MB/s, and compare random with sequential patterns. Explain any rise in bandwidth accompanied by lower IOPS using the fixed-size conversion. Discuss repeat variability rather than attributing every small difference to a device effect.]")
    r.p("Maximum observed random 4 KiB IOPS  Read [TBD]  Write [TBD]\nMaximum observed sequential throughput  Read [TBD] MiB/s  Write [TBD] MiB/s\nCase identifiers and fixed knobs for each maximum  [TBD]")
    r.p("Use analysis/maxima.csv to identify the maximum within the tested matrix. This is an observed profile maximum, not a manufacturer specification or proof of the device's absolute capability.")

    r.page(); r.h("6 Mix queue depth and worker results")
    r.fig("04_mix","Figure 4 Read/write mix at fixed QD and block size within each pattern")
    r.p("[TBD Compare 100/0, 70/30, 50/50 and 0/100. Use achieved read fractions and separate read/write rows to explain any asymmetric latency effect. Evaluate possible controller or shared-media contention with evidence.]")
    r.fig("05_queue_tradeoff","Figure 5 Single throughput versus mean total latency curve with QD labels for all four workloads")
    r.fig("06_queue_scaling","Figure 6 Throughput and mean latency as configured QD increases")
    knee_rows=r.csv("knees.csv")
    table=[]
    for pattern in ("random","sequential"):
        for read in (100,0):
            row=next((v for v in knee_rows if v["pattern"]==pattern and int(v["read_pct"])==read),None)
            table.append([pattern+" "+("read" if read==100 else "write"),row["selected_qd"] if row else "TBD",("Candidate" if row["plateau_observed"]=="True" else "Not established") if row else "TBD","TBD"])
    r.table(["Workload","Selected QD","Plateau assessment","Supporting throughput and latency"],table,[1.65,.85,1.4,3.1])
    r.p("[TBD Accept or qualify the heuristic candidate using adjacent points and SD. If no plateau is observed, state that the knee was not established in this range and identify the fallback depth. Explain whether the plateau could reflect host CPU, virtualization, interface or media limits; do not choose a cause without evidence.]")
    r.fig("07_little_law","Figure 7 IOPS times mean completion latency compared with configured aggregate QD")
    r.p("Numeric Little's Law check  Workload and QD [TBD]; lambda [TBD] requests/s; W [TBD] us / 1000000; L [TBD] requests; configured QD [TBD]. Compare submitted and total-latency estimates and discuss underfilling, submission overhead, or nonstationarity. Use analysis/little_law.csv and per-job depth buckets.")
    r.fig("08_parallel","Figure 8 Worker count at fixed aggregate QD and fixed total address range")
    r.p("[TBD Compare one deep queue with several shallower workers at aggregate QD 16 and 64. Relate changes to CPU overhead or submission parallelism while checking achieved occupancy.]")

    r.page(); r.h("7 Tail latency working set and device state")
    tailrows=[]
    for pattern in ("random","sequential"):
        for read in (100,0):
            available=[x for x in r.rows if x["family"]=="tail" and x["direction"]=="total" and x["pattern"]==pattern and int(x["read_pct"])==read]
            for role in ("mid","comparison"):
                row=next((x for x in available if (x["role"]=="mid" if role=="mid" else x["role"]!="mid")),None)
                tailrows.append([f"{pattern}\n{'R' if read==100 else 'W'} {row['role'] if row else role}",row["iodepth"] if row else "TBD",*(r.number(row,m) for m in ("p50_us","p95_us","p99_us","p999_us")),str(int(float(row.get("clat_samples_mean", row["io_count_mean"])))) if row else "TBD"])
    r.p("Completion percentiles below are microseconds and are mean per-trial percentile values. N is the mean completed-I/O count per trial; include its range and SD when interpreting a sparse tail. The comparison depth is near-knee only when the declared rule has evidence for a plateau.")
    r.table(["Workload and role","QD","p50","p95","p99","p99.9","Mean N"],tailrows,[2.05,.55,.8,.8,.8,.9,1.1],font=8.7)
    r.fig("09_tails","Figure 9 Completion-latency percentiles at middle and selected comparison QD")
    r.fig("10_tail_survival","Figure 10 Labeled pooled histogram survival curves across trials")
    r.p("[TBD Compare p99.9/p50 and p99/p50 across depths. State trial counts and any small-N diagnostics. Explain whether increased concurrency mainly changes central latency, tails, or both, and what that means for response-time-sensitive workloads.]")
    r.h("Working set and observed write state",2)
    r.fig("11_working_set","Figure 11 Guest address-range effects at random 4 KiB and QD 16")
    r.fig("12_state","Figure 12 Throughput and mean completion latency during each sustained write window")
    r.p("[TBD Record tested ranges, compare IOPS and p99, and explain that guest address windows do not select physical flash LBAs. For each state trial, compare first and last thirds using state_windows.csv. Pair changes with available SSD temperature or host reliability observations. If no change occurs, limit the conclusion to this duration and range.]")

    r.page(); r.h("8 Misleading setup and limitations")
    r.h("Warm buffered reads as an SSD benchmark pitfall",2)
    r.p("The controlled pitfall experiment uses a finite buffered sequential read to warm the same initialized 256 MiB or smaller range before each case. Both measured cases use random 4 KiB reads, psync, QD 1, one worker, and the same range. Only direct=1 versus direct=0 changes. Holding the engine constant avoids confusing the cache effect with asynchronous submission overhead.")
    r.fig("13_pitfall","Figure 13 Direct and warmed buffered reads under the same synchronous workload")
    r.p("Buffered reads can be satisfied from guest RAM, so presenting their request rate as physical SSD IOPS can be misleading. A cache benefit depends on effective cache residency and the storage layers involved. If the two controls are similar, check whether warming completed, whether the range fits available memory, and whether host/virtualization overhead dominates. No global page-cache flush is performed. direct=1 concerns the guest cache and cannot certify the behavior of every host cache.")
    r.p("[TBD Report each control's IOPS and latency with trial variation, compute buffered/direct IOPS, and explain the observed direction and size of the effect. State whether the experiment demonstrated an effect or only bounded it.]")
    r.h("Storage path and durability",2)
    r.p("WSL2 stores a distribution in an ext4 virtual hard disk [3]. This experiment includes the Linux filesystem, virtual disk and host storage layers. Linux-local storage avoids the separate Windows-mounted file path that Microsoft advises against for Linux performance work [2], but the results still describe a virtualized path. Physical SSD temperature, interface detail and hardware queue depth may not be exposed to the guest. Guest free space may reflect the VHD limit rather than actual remaining host capacity [3].")
    r.p("Write completion is measured without per-I/O fsync or a power-loss test. Consequently these write latencies do not establish durable NAND commit latency. The full-file initialization does not reset physical drive history or force a standardized steady state. Address ranges within a file are mediated by filesystem allocation, VHDX mapping and the flash translation layer.")
    r.h("Attribution and uncertainty",2)
    r.p("SLC caching, garbage collection, thermal throttling, host caching, background I/O and CPU limits can produce similar changes. A throughput drop alone does not distinguish them. Use measured temperature and chronology where available, and label other explanations as hypotheses. The final part of a 120-second write window is an observed late window, not automatically a steady state. Shared drive history also limits independence between repeated trials.")
    r.p("[TBD List unavailable sensors, unexplained anomalies, incomplete settings, interruptions, CPU or memory pressure, and the limitations that materially affect the final conclusion.]")

    r.page(); r.h("9 Conclusions and submission evidence")
    r.p("[TBD Summarize the measured QD 1 latency, maximum observed random IOPS and sequential bandwidth, knee or lack of one, and tail-latency tradeoff. Explain which tested settings best serve latency-sensitive and throughput-oriented work. State the strongest supported device-state observation and the effect of the misleading setup. Keep every numerical claim traceable to a case and acknowledge the WSL scope.]")
    r.h("Rubric coverage",2)
    r.table(["Category","Points","Evidence"],[
        ["Safety and QD 1 baselines",30,"Owned-file guards, initialization evidence, metadata and Section 5"],
        ["Block size and patterns",35,"Complete 24-case matrix; Figures 2 and 3; Section 5"],
        ["Mix and queue analysis",40,"Mix/QD/worker sweeps, knee evidence and numeric Little's Law check; Section 6"],
        ["Tails and device state",25,"p50/p95/p99/p99.9, N, ranges, state chronology and temperature limits; Section 7"],
        ["Reproducibility and hygiene",25,"Configs, raw data, hashes, repeated trials, environment, diagnostics and coverage"],
        ["AI and pitfall analysis",15,"Disclosure below and controlled cache experiment; Section 8"],
        ["Report quality",15,"Labeled figures, units, concise supported interpretation and limitations"],
        ["Total",185,"Final empirical evidence must accompany the implementation"]
    ],[2.0,.55,4.45],font=8.8)
    r.h("AI use disclosure",2)
    r.p("OpenAI ChatGPT assisted in generating the experiment code, parsing and plotting routines, documentation, and report structure from the supplied assignment. The author is responsible for running the target-laptop experiments, verifying raw results and generated statements, and writing the final empirical interpretation. No target-laptop performance results are supplied in the starter report. Update this disclosure to reflect any additional assistance.")
    r.h("References",2)
    refs=[
        "[1] fio project. Official fio documentation. https://fio.readthedocs.io/en/latest/fio_doc.html",
        "[2] Microsoft. Working across Windows and Linux file systems. https://learn.microsoft.com/en-us/windows/wsl/filesystems",
        "[3] Microsoft. How to manage WSL disk space. https://learn.microsoft.com/en-us/windows/wsl/disk-space",
        "[4] Microsoft. Get PhysicalDisk. https://learn.microsoft.com/en-us/powershell/module/storage/get-physicaldisk",
        "[5] Microsoft. Get StorageReliabilityCounter. https://learn.microsoft.com/en-us/powershell/module/storage/get-storagereliabilitycounter",
        "[6] Advanced Computer Systems. Project 3 SSD Performance Profiling assignment and 185-point grading rubric supplied with the project."
    ]
    for text in refs:
        p=r.p(text)
        for run in p.runs: run.font.size=Pt(8)
        p.paragraph_format.space_after=Pt(3)
    output=Path(output)
    output.parent.mkdir(parents=True,exist_ok=True)
    if output.exists() or output.with_suffix(".md").exists():
        raise ValueError("Choose a new output path; existing reports are never overwritten")
    d.save(output)
    output.with_suffix(".md").write_text("\n".join(r.md),encoding="utf-8")
    print(f"Saved {output} and matching Markdown. Empirical interpretation fields remain TBD.")


if __name__=="__main__":
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument("--results",help="Existing results directory containing analysis/ and optional plots/")
    p.add_argument("--output",required=True)
    a=p.parse_args()
    build(a.output,a.results)
'''

def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--results',default='results/full_01')
    args=ap.parse_args()
    root=Path.cwd().resolve()
    if not (root/'ssdbench'/'runner.py').is_file():
        raise ValueError('First cd to ~/ssd_profiling, then run this script.')
    sys.path.insert(0,str(root))
    from ssdbench.common import digest,dump_json,read_json,utcnow
    from ssdbench.runner import journal_events,source_hashes
    from ssdbench.safety import locked,target_directory,validate_target
    out=Path(args.results).expanduser().resolve()
    manifest_path=out/'manifest.json'
    m=read_json(manifest_path)
    directory=target_directory(m['target_directory'])
    with locked(directory):
        target=validate_target(directory)
        if target['uuid']!=m['target']['uuid']:
            raise ValueError('Target identity differs from the saved run. Nothing changed.')
        current=source_hashes(root)
        def fingerprint(sources):
            return digest({'parameters':m['parameters'],'target_uuid':m['target']['uuid'],
                           'source':sources,'fio_version':m['fio_version']})
        if current!=m['source_sha256'] or fingerprint(current)!=m['fingerprint']:
            raise ValueError('Current code differs from the saved checkpoint. No files were changed; retain your backups and request help.')
        if all(current.get(name)==sha for name,sha in NEW_HASHES.items()):
            print('This repair is already applied; the checkpoint matches the repaired code.')
            return
        for name,sha in OLD_HASHES.items():
            if current.get(name)!=sha:
                raise ValueError('Unexpected version of '+name+'. No files were changed.')
        events=journal_events(out)
        completed=[e for e in events if e.get('status')=='complete' and e.get('phase')=='measured']
        # Validate retained bytes before changing checkpoint provenance.
        for e in completed:
            path=(out/e['raw']).resolve()
            if out not in path.parents:
                raise ValueError('Raw path leaves results directory.')
            if hashlib.sha256(path.read_bytes()).hexdigest()!=e['raw_sha256']:
                raise ValueError('A retained raw result changed: '+str(path))
        namespace={'__name__':'ssdbench._histogram_fix_preview','__package__':'ssdbench'}
        exec(compile(PAYLOADS['ssdbench/analysis.py'],'analysis.py','exec'),namespace)
        parse=namespace['parse_document']
        for e in completed:
            parse(read_json(out/e['raw']),e['case'],e['run_id'],e['trial'])
        histogram_failures=[e for e in events if e.get('status')=='failed'
                            and 'Completion histogram count differs from completed I/O count' in e.get('error','')]
        for e in histogram_failures:
            parse(read_json(out/e['raw']),e['case'],e['run_id'],e['trial'])
        old_bytes={name:(root/name).read_bytes() for name in PAYLOADS}
        manifest_bytes=manifest_path.read_bytes()
        backup_manifest=out/'manifest.before_histogram_fix.json'
        if backup_manifest.exists():
            raise ValueError('A repair backup already exists. No files changed; request help before retrying.')
        backups={name:(root/name).with_name(Path(name).name+'.before_histogram_fix') for name in PAYLOADS}
        if any(p.exists() for p in backups.values()):
            raise ValueError('Source repair backups already exist. No files changed.')
        backup_manifest.write_bytes(manifest_bytes)
        for name,p in backups.items():
            p.write_bytes(old_bytes[name])
        try:
            for name,text in PAYLOADS.items():
                destination=root/name
                temp=destination.with_suffix('.repair_tmp')
                temp.write_text(text,encoding='utf-8')
                temp.replace(destination)
            updated=source_hashes(root)
            changed={name for name in set(current)|set(updated) if current.get(name)!=updated.get(name)}
            if changed!=set(PAYLOADS):
                raise ValueError('Unexpected code changes during repair.')
            m.setdefault('compatibility_updates',[]).append({
                'utc':utcnow(),
                'reason':'Validate histogram against clat N; retain bounded asynchronous ramp-boundary observations and distinct I/O/sample counts.',
                'previous_fingerprint':m['fingerprint'],
                'previous_source_sha256':current,
                'updated_source_sha256':updated,
                'manifest_backup':backup_manifest.name,
                'source_backups':{name:str(p.relative_to(root)) for name,p in backups.items()},
                'completed_runs_preserved':len(completed),
                'histogram_failures_validated':len(histogram_failures)})
            m['source_sha256']=updated
            m['fingerprint']=fingerprint(updated)
            dump_json(manifest_path,m)
        except BaseException:
            for name,blob in old_bytes.items():
                (root/name).write_bytes(blob)
            manifest_path.write_bytes(manifest_bytes)
            raise
    print('Parser repaired. Original code and manifest backed up; raw results unchanged.')
    print('Completed measured runs preserved:',len(completed))
    print('The failed parser attempt remains archived and will be retried on resume.')
    subprocess.run([sys.executable,str(root/'scripts'/'ssdbench.py'),'analyze','--results',str(out)],check=True)
    print('Resume your original full command with --resume. Do not initialize again.')

if __name__=='__main__':
    try:
        main()
    except (ValueError,OSError,KeyError,subprocess.SubprocessError) as exc:
        print('Repair stopped:',exc,file=sys.stderr)
        raise SystemExit(2)
